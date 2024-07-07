from flask import Flask, request, jsonify, send_file
import subprocess
import pandas as pd
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
import os

app = Flask(__name__)

def get_subdomains(domain):
    command = f"assetfinder {domain}"
    result = subprocess.run(command, shell=True, capture_output=True, text=True)
    subdomains = result.stdout.split()
    return subdomains

def classify_subdomain(subdomain):
    keywords = {
        "files": ["file", "storage", "docs"],
        "backup": ["backup", "bkp", "archive"],
        "conf": ["config", "setup", "settings"],
        "int": ["internal", "intranet"],
        "password": ["auth", "login", "secure"],
        "auth": ["oauth", "authentication"],
        "acl": ["access", "acl"],
        "log": ["log", "report", "audit"],
        "source": ["dev", "source", "src", "code"],
        "exe": ["app", "exe", "bin"],
        "test": ["test", "devtest", "testing"]
    }
    for key, words in keywords.items():
        if any(word in subdomain for word in words):
            return key
    return "int"

def classify_subdomains(subdomains):
    data = []
    for i, subdomain in enumerate(subdomains):
        asset_type = classify_subdomain(subdomain)
        data.append({
            "ID activo": str(i + 1),  # Asegurarse de que sea string
            "Tipo de activo": asset_type,
            "Nombre del activo": subdomain,
            "Descripción": f"Subdominio detectado por Assetfinder clasificado como {asset_type}"
        })
    print(data)  # Imprimir los datos para depuración
    return data


def save_to_excel(data):
    df = pd.DataFrame(data)
    # Convertir nuevamente los nombres de columnas a strings
    df.columns = [str(column) for column in df.columns]
    print(df.head())  # Imprimir las primeras filas para depuración
    file_path = "activos_digitales.xlsx"
    with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Activos Digitales")
        workbook = writer.book
        worksheet = writer.sheets["Activos Digitales"]
        for column_cells in worksheet.columns:
            length = max(len(str(cell.value)) for cell in column_cells)
            worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = length
        table = Table(displayName="TablaActivos", ref=worksheet.dimensions)
        style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                               showLastColumn=False, showRowStripes=True, showColumnStripes=True)
        table.tableStyleInfo = style
        worksheet.add_table(table)
    return file_path



@app.route('/search', methods=['GET'])
def search():
    domain = request.args.get('domain')
    if not domain:
        return jsonify({"message": "No domain provided"}), 400
    subdomains = get_subdomains(domain)
    data = classify_subdomains(subdomains)
    file_path = save_to_excel(data)
    return jsonify({"message": "Search completed", "data": data, "file_path": file_path})

@app.route('/download', methods=['GET'])
def download():
    file_path = "activos_digitales.xlsx"
    if os.path.exists(file_path):
        return send_file(file_path, as_attachment=True)
    else:
        return jsonify({"message": "File not found"}), 404

if __name__ == '__main__':
    app.run(debug=True)
